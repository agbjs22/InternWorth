from modules.pay_extractor import extract_amount, score as pay_score
from modules.red_flag_detector import score as red_flag_score
from scoring import score_posting
from modules.file_extractor import extract_file_text, posting_from_text
from docx import Document
import io


def test_extracts_indian_stipend_formats():
    assert extract_amount("₹15,000 / month") == 15000
    assert extract_amount("Rs. 12k per month") == 12000
    assert extract_amount("Unpaid internship") == 0


def test_pay_handles_missing_stipend_explainably():
    result = pay_score("", "Pune")
    assert result["sub_score"] == 35
    assert result["note"]


def test_red_flags_reduce_score():
    result = red_flag_score("Unpaid role. Registration fee required. 12 hours daily.")
    assert result["sub_score"] < 50
    assert "unpaid" in result["flags"]


def test_full_score_preserves_explainability_contract():
    result = score_posting({
        "title": "Software Intern", "company": "Example Labs", "location": "Remote",
        "stipend": "₹25,000/month", "company_info": "Founded in 2018. Website and clients listed.",
        "description": "Build React dashboard features, write tests, use Git and APIs. Mentorship and a full-time offer are available.",
    })
    assert 0 <= result["worth_score"] <= 100
    assert len(result["breakdown"]) == 6
    assert all("sub_score" in item and "note" in item for item in result["breakdown"].values())


def test_extracts_docx_listing_and_builds_posting_hints():
    document = Document()
    document.add_paragraph("Data Analyst Intern")
    document.add_paragraph("Company: Example Labs")
    document.add_paragraph("Location: Remote")
    document.add_paragraph("Stipend: ₹20,000 per month")
    document.add_paragraph("Build dashboards using Python and SQL with a mentor.")
    buffer = io.BytesIO()
    document.save(buffer)
    text = extract_file_text("listing.docx", buffer.getvalue())
    posting = posting_from_text(text)
    assert "dashboards" in text
    assert posting["company"] == "Example Labs"
    assert posting["stipend"] == "₹20,000 per month"
