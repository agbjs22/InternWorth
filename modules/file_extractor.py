"""In-memory extraction for a single user-uploaded listing file."""
from __future__ import annotations

import io
import re
import shutil
from pathlib import Path

import pymupdf
from docx import Document
from PIL import Image
import pytesseract

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".png", ".jpg", ".jpeg", ".webp"}

if not shutil.which("tesseract"):
    for executable in (Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe"), Path(r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe")):
        if executable.exists():
            pytesseract.pytesseract.tesseract_cmd = str(executable)
            break


class FileExtractionError(Exception):
    pass


def _ocr_image(image: Image.Image) -> str:
    try:
        return pytesseract.image_to_string(image)
    except pytesseract.TesseractNotFoundError as error:
        raise FileExtractionError("Screenshot OCR needs Tesseract installed on the server. You can still upload a PDF or DOCX, or use manual entry.") from error


def extract_file_text(filename: str, content: bytes) -> str:
    extension = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension not in SUPPORTED_EXTENSIONS:
        raise FileExtractionError("Use a PDF, Word .docx, PNG, JPG, or WEBP file.")
    if not content:
        raise FileExtractionError("The uploaded file is empty.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise FileExtractionError("Please upload a file smaller than 10 MB.")
    try:
        if extension == ".docx":
            document = Document(io.BytesIO(content))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            for table in document.tables:
                text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        elif extension == ".pdf":
            pdf = pymupdf.open(stream=content, filetype="pdf")
            text = "\n".join(page.get_text("text") for page in pdf)
            if len(text.strip()) < 30:
                text = "\n".join(_ocr_image(Image.open(io.BytesIO(page.get_pixmap(matrix=pymupdf.Matrix(2, 2)).tobytes("png")))) for page in pdf)
        else:
            text = _ocr_image(Image.open(io.BytesIO(content)))
    except FileExtractionError:
        raise
    except Exception as error:
        raise FileExtractionError("We could not read this file. Try a clearer image, text-based PDF, or .docx document.") from error
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 20:
        raise FileExtractionError("Not enough internship text could be detected. Please try manual entry.")
    return text[:30000]


def posting_from_text(text: str) -> dict[str, str]:
    """Conservative field hints; full text remains the scoreable description."""
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    compact = " ".join(lines)
    def labelled(label: str) -> str:
        match = re.search(rf"\b{label}\s*[:\-]\s*([^\n|]{{2,100}})", text, re.IGNORECASE)
        return match.group(1).strip() if match else ""
    stipend = labelled("(?:stipend|salary|compensation|pay)")
    if not stipend:
        amount = re.search(r"(?:₹|rs\.?|inr)\s*[\d,.]+\s*(?:k|l)?(?:\s*(?:/|per)\s*(?:month|month|mo))?", compact, re.IGNORECASE)
        stipend = amount.group(0) if amount else ""
    return {
        "title": labelled("(?:role|position|job title|internship)") or (lines[0] if lines else "Internship opportunity"),
        "company": labelled("(?:company|organization|employer)"),
        "location": labelled("(?:location|work location|job location)"),
        "stipend": stipend,
        "description": text,
        "company_info": "",
    }
